from pathlib import Path
import csv, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
RUN=ROOT/'air_quality_forecasting/outputs/runs/current'
OUT=ROOT/'reports'
metrics=list(csv.DictReader((RUN/'metrics.csv').open()))
city=list(csv.DictReader((RUN/'city_metrics.csv').open()))
sel=json.loads((RUN/'selection.json').read_text())
refs=[
'Open-Meteo, “Air Quality API,” https://open-meteo.com/en/docs/air-quality-api (accessed Oct. 20, 2026).',
'Open-Meteo, “Historical Weather API,” https://open-meteo.com/en/docs/historical-weather-api (accessed Oct. 20, 2026).',
'scikit-learn developers, “Linear Models,” https://scikit-learn.org/stable/modules/linear_model.html (accessed Oct. 20, 2026). Implementation version: 1.7.2.',
'R. J. Hyndman and G. Athanasopoulos, Forecasting: Principles and Practice, 3rd ed., OTexts, 2021, “The forecaster’s toolbox,” https://otexts.com/fpp3/toolbox.html (accessed Oct. 20, 2026).',
'Air quality forecasting project, local experiment artifacts, outputs/runs/current: manifest.json, config_used.yaml, selection.json, validation_scores.csv, metrics.csv, city_metrics.csv and predictions.csv, Oct. 20, 2026. Base repository: https://github.com/phanduyhoang22/air_quality_forecasting. Results use the locally modified working tree, not the unmodified base commit.'
]
# Paired prose keeps scientific claims identical across the two editions.
pages=[
[
('title','Direct PM2.5 Forecasting with Linear and Polynomial Regression','Dự báo trực tiếp PM2.5 bằng hồi quy tuyến tính và đa thức'),
('sub','A comparative study across four Vietnamese cities at one three and six hours','Nghiên cứu so sánh tại bốn thành phố Việt Nam với các mốc một ba và sáu giờ'),
('h','Abstract','Tóm tắt'),
('p','This study evaluates direct PM2.5 concentration forecasts one, three and six hours ahead in Hanoi, Hai Phong, Da Nang and Ho Chi Minh City. Simple linear, multiple linear, polynomial and Ridge regression are compared with a persistence baseline using hourly API data covering 2023–2024. Causal feature construction, chronological expanding validation and a six-hour boundary purge separate model selection from the final test. Of 70,176 city-hour records, 52,008 meet the common feature and target requirements. Multiple linear regression, tied with Ridge at zero regularization, is selected by validation at every horizon. On 7,804 test records per horizon, its RMSE is 3.908, 9.430 and 13.974 µg/m³, respectively, reducing pooled RMSE by 19.41%, 18.30% and 20.29% relative to persistence. These aggregate gains do not generalize uniformly: Da Nang performs worse than persistence at all three horizons. Missing meteorological data, model-derived pollution inputs and unverified real-time availability restrict the conclusions to a retrospective benchmark. The results support a transparent regression baseline rather than a validated operational warning system.',
'Nghiên cứu đánh giá dự báo trực tiếp nồng độ PM2.5 trước một, ba và sáu giờ tại Hà Nội, Hải Phòng, Đà Nẵng và TP. Hồ Chí Minh. Hồi quy tuyến tính đơn, tuyến tính đa biến, đa thức và Ridge được so sánh với baseline persistence trên dữ liệu API theo giờ giai đoạn 2023–2024. Đặc trưng chỉ dùng thông tin đến thời điểm dự báo; validation mở rộng theo thời gian và khoảng loại trừ sáu giờ tách lựa chọn mô hình khỏi tập test cuối. Trong 70.176 bản ghi thành phố–giờ, 52.008 bản ghi đáp ứng yêu cầu chung về đặc trưng và nhãn. Validation chọn hồi quy tuyến tính đa biến tại cả ba mốc, đồng hạng với Ridge khi hệ số phạt bằng không. Trên 7.804 bản ghi test mỗi mốc, RMSE lần lượt đạt 3,908; 9,430 và 13,974 µg/m³, giảm 19,41%; 18,30% và 20,29% so với persistence. Tuy nhiên, mô hình kém persistence tại Đà Nẵng ở cả ba mốc. Dữ liệu khí tượng thiếu, nguồn ô nhiễm dựa trên mô hình và giả định chưa kiểm chứng về thời điểm dữ liệu sẵn có giới hạn kết luận trong phạm vi đánh giá hồi cứu. Kết quả cung cấp một hệ thống hồi quy minh bạch, chưa xác nhận khả năng cảnh báo vận hành thực tế.'),
('p','Keywords: PM2.5; multiple linear regression; polynomial regression; Ridge; time series validation; persistence.','Từ khóa: PM2.5; hồi quy tuyến tính đa biến; hồi quy đa thức; Ridge; validation chuỗi thời gian; persistence.'),
('h','1 Introduction','1 Giới thiệu'),
('p','Short-horizon concentration forecasting is a useful setting for studying regression because the target is numerical, observations are ordered in time, and a strong baseline is immediately available. The research question is whether meteorology, temporal features and recent concentration history improve predictions over carrying the latest available concentration forward. A second question concerns how that improvement changes across horizons and cities. The study does not estimate causal effects of weather on pollution.',
'Dự báo nồng độ ngắn hạn là bài toán phù hợp để nghiên cứu hồi quy vì biến đích là số liên tục, quan sát có thứ tự thời gian và có thể xây dựng baseline rõ ràng. Câu hỏi nghiên cứu là liệu khí tượng, đặc trưng thời gian và lịch sử nồng độ gần nhất có cải thiện dự báo so với giữ nguyên nồng độ hiện tại hay không. Câu hỏi thứ hai là mức cải thiện thay đổi thế nào theo thành phố và khoảng dự báo. Nghiên cứu không ước lượng tác động nhân quả của thời tiết lên ô nhiễm.'),
],
[
('h','2 Data and forecasting task','2 Dữ liệu và bài toán dự báo'),
('p','The snapshot contains 17,544 hourly rows per city from 1 January 2023 00:00 to 31 December 2024 23:00, in Asia/Ho_Chi_Minh time (UTC+7). PM2.5 comes from the Open-Meteo air-quality service, while weather variables come from its historical weather service [1], [2]. These are gridded model and historical weather products, not an independently collected station-monitoring dataset. The API snapshot is preserved with source metadata and SHA-256 hashes [5].',
'Snapshot gồm 17.544 dòng theo giờ cho mỗi thành phố từ 00:00 ngày 01/01/2023 đến 23:00 ngày 31/12/2024, múi giờ Asia/Ho_Chi_Minh (UTC+7). PM2.5 lấy từ dịch vụ chất lượng không khí Open-Meteo; biến khí tượng lấy từ dịch vụ thời tiết lịch sử [1], [2]. Đây là dữ liệu lưới từ mô hình và sản phẩm thời tiết lịch sử, không phải bộ số đo trạm độc lập. Snapshot API được lưu cùng metadata nguồn và mã băm SHA-256 [5].'),
('cities',),
('p','At forecast origin t, after observations for hour t are assumed available, a separate model predicts PM2.5 at t+h for h ∈ {1, 3, 6}. The targets are concentrations at individual future hours, not averages over the next h hours. Models share observations across cities and include city indicators. Each horizon is fitted directly; intermediate future concentrations are never supplied as predictors.',
'Tại thời điểm phát hành t, sau khi giả định dữ liệu giờ t đã sẵn có, một mô hình riêng dự báo PM2.5 tại t+h với h ∈ {1, 3, 6}. Nhãn là nồng độ tại từng giờ tương lai, không phải trung bình trong h giờ tiếp theo. Mô hình dùng chung dữ liệu các thành phố và có biến chỉ báo thành phố. Mỗi mốc được học trực tiếp; nồng độ thật tại các giờ tương lai trung gian không được dùng làm đầu vào.'),
('h','2 1 Missing data and eligible observations','2 1 Dữ liệu thiếu và mẫu đủ điều kiện'),
('p','PM2.5 has no missing values in this snapshot. Boundary-layer height is missing for 4,368 hours per city, from 1 January 2024 07:00 through 1 July 2024 06:00. The pipeline restores the hourly index and retains missing values without interpolation or backward filling. Lagged features, rolling windows and all three targets must be complete. The resulting common sample has 52,008 of 70,176 rows (74.11%); 18,168 rows are excluded. Counts also reflect the initial lag history and final unavailable future labels, not only missing weather.',
'PM2.5 không có giá trị thiếu trong snapshot này. Độ cao lớp biên khí quyển thiếu 4.368 giờ mỗi thành phố, từ 07:00 ngày 01/01/2024 đến 06:00 ngày 01/07/2024. Pipeline khôi phục trục giờ đều và giữ giá trị thiếu, không nội suy hay điền ngược. Các độ trễ, cửa sổ trượt và cả ba nhãn phải đầy đủ. Tập mẫu chung còn 52.008 trên 70.176 dòng (74,11%); 18.168 dòng bị loại. Số dòng loại còn bao gồm phần lịch sử ban đầu chưa đủ độ trễ và phần cuối chưa có nhãn tương lai, không chỉ do khí tượng thiếu.'),
('p','Complete-case filtering changes the temporal coverage of the experiment. Consequently, the split proportions refer to eligible timestamps, and the reported accuracy must not be extended to missing periods. No local imputation does not establish that the upstream provider performed no interpolation or retrospective processing.',
'Lọc mẫu đầy đủ làm thay đổi độ phủ thời gian của thí nghiệm. Vì vậy, tỷ lệ chia tập áp dụng lên các timestamp đủ điều kiện và không thể suy rộng độ chính xác sang các khoảng bị thiếu. Không nội suy trong mã nguồn không có nghĩa nhà cung cấp dữ liệu không nội suy hoặc xử lý hồi cứu ở phía nguồn.'),
],
[
('h','3 Regression methods and predictors','3 Phương pháp hồi quy và đặc trưng'),
('p','The persistence baseline predicts the latest known concentration for every horizon. Simple linear regression learns an intercept and one coefficient on current PM2.5. Multiple linear regression uses ordinary least squares on 42 explicitly allowed features. Polynomial regression expands four variables—current PM2.5, temperature, relative humidity and wind speed—to degree two or three, including interactions, while retaining other features linearly. Ridge uses the same feature set as multiple linear regression with an L2 coefficient penalty [3].',
'Baseline persistence dự báo bằng nồng độ gần nhất đã biết cho mọi mốc. Hồi quy tuyến tính đơn học hệ số chặn và một hệ số của PM2.5 hiện tại. Hồi quy tuyến tính đa biến dùng bình phương tối thiểu trên 42 đặc trưng được khai báo tường minh. Hồi quy đa thức mở rộng bốn biến—PM2.5 hiện tại, nhiệt độ, độ ẩm tương đối và tốc độ gió—đến bậc hai hoặc ba, có tương tác; các biến còn lại giữ tuyến tính. Ridge dùng cùng bộ đặc trưng với hồi quy đa biến và phạt L2 lên hệ số [3].'),
('p','The 42 predictors comprise nine weather variables; current PM2.5; six cyclic calendar encodings; a weekend indicator; five derived weather terms; nine PM2.5 lags at 1, 2, 3, 6, 12, 24, 48, 72 and 168 hours; eight rolling statistics; and three city indicators, with Hanoi as reference. Rolling means and standard deviations use windows of 3, 6, 12 and 24 hours. They include hour t, remain within each city, and require complete windows. Derived terms include temperature–humidity interaction, wind components, dew-point depression and pressure change.',
'Có 42 đầu vào: chín biến khí tượng; PM2.5 hiện tại; sáu mã hóa chu kỳ lịch; chỉ báo cuối tuần; năm biến khí tượng suy dẫn; chín độ trễ PM2.5 tại 1, 2, 3, 6, 12, 24, 48, 72 và 168 giờ; tám thống kê trượt; ba chỉ báo thành phố với Hà Nội làm tham chiếu. Trung bình và độ lệch chuẩn trượt dùng cửa sổ 3, 6, 12 và 24 giờ, bao gồm giờ t, tính riêng từng thành phố và yêu cầu cửa sổ đầy đủ. Biến suy dẫn gồm tương tác nhiệt độ–độ ẩm, các thành phần gió, chênh lệch nhiệt độ với điểm sương và biến thiên áp suất.'),
('h','3 1 Optimization and preprocessing','3 1 Tối ưu hóa và tiền xử lý'),
('p','For a fixed horizon, the fitted prediction is an intercept plus a weighted sum of the transformed inputs. Ordinary least squares minimizes mean squared residuals. The Ridge objective adds alpha times the sum of squared slope coefficients; the intercept is not penalized. The experimental solver uses a centered NumPy closed-form solution. At alpha = 0 it uses least squares, allowing rank-deficient designs. A gradient-descent implementation is retained and tested against the closed-form solution; it is not the solver used to produce these results.',
'Với mỗi mốc, dự đoán là hệ số chặn cộng tổng có trọng số của các đầu vào đã biến đổi. Bình phương tối thiểu giảm thiểu trung bình bình phương phần dư. Hàm mục tiêu Ridge cộng thêm alpha nhân tổng bình phương hệ số góc; hệ số chặn không bị phạt. Thí nghiệm dùng nghiệm trực tiếp có định tâm bằng NumPy. Khi alpha = 0, bộ giải dùng bình phương tối thiểu, xử lý được thiết kế thiếu hạng. Cài đặt gradient descent được giữ và kiểm thử đối chiếu với nghiệm trực tiếp, nhưng không dùng để tạo kết quả báo cáo.'),
('p','Standardization is fitted on each training fold only, then reused unchanged on validation. The final transforms are refitted on development data before testing. Targets remain in µg/m³. The Ridge search grid is {0, 0.001, 0.01, 0.1, 1}; polynomial degrees are {2, 3}. Because the custom objective uses mean squared error, its alpha corresponds to n_train × alpha in the sum-of-squares convention used by scikit-learn Ridge. The design follows regression, regularization, resampling and validation-based model-selection concepts; chronological validation adapts those concepts to forecasting.',
'Chuẩn hóa chỉ được fit trên train của từng fold rồi áp dụng nguyên trạng lên validation. Trước test, phép biến đổi được fit lại trên development. Nhãn giữ đơn vị µg/m³. Lưới alpha Ridge là {0; 0,001; 0,01; 0,1; 1}; bậc đa thức là {2; 3}. Vì hàm mục tiêu tự cài đặt dùng trung bình bình phương sai số, alpha tương ứng trong quy ước tổng bình phương của scikit-learn Ridge là n_train × alpha. Thiết kế bám các khái niệm hồi quy, regularization, resampling và chọn mô hình bằng validation; chia theo thời gian là sự điều chỉnh phù hợp cho dự báo.'),
],
[
('h','4 Experimental protocol','4 Thiết kế thực nghiệm'),
('p','Unique eligible timestamps are split chronologically: the first 70% initialize training, the next 15% provide three expanding validation windows, and the final 15% form the held-out test. All cities at the same timestamp remain on the same side of a boundary. Training rows whose six-hour target reaches or exceeds the next validation or test boundary are removed. Validation labels also remain inside their respective windows. This procedure preserves chronological ordering and follows the rationale of rolling-origin evaluation [4].',
'Các timestamp đủ điều kiện được chia theo thứ tự thời gian: 70% đầu làm train ban đầu, 15% tiếp theo tạo ba cửa sổ validation mở rộng, 15% cuối làm test giữ riêng. Các thành phố tại cùng timestamp luôn ở cùng phía ranh giới. Mẫu train có nhãn sáu giờ chạm hoặc vượt đầu validation hay test bị loại. Nhãn validation cũng phải nằm trong cửa sổ tương ứng. Quy trình giữ trật tự thời gian và phù hợp với nguyên tắc đánh giá theo thời điểm dự báo cuốn chiếu [4].'),
('folds',),
('p','Validation starts at 22 July 2024 05:00. The final test origins run from 11 October 2024 11:00 through 31 December 2024 17:00. After purging, development contains 44,180 rows; test contains 7,804 rows, or 1,951 per city, for every method and horizon. The shared six-hour purge is conservative for the shorter horizons. Hyperparameters and the winning family are selected by pooled validation RMSE, computed from pooled squared residuals, not the mean of fold RMSEs. Each family is then refitted on development and evaluated on the same test rows [5].',
'Validation bắt đầu lúc 05:00 ngày 22/07/2024. Các thời điểm phát hành trong test trải từ 11:00 ngày 11/10/2024 đến 17:00 ngày 31/12/2024. Sau loại trừ ranh giới, development có 44.180 dòng; test có 7.804 dòng, tức 1.951 dòng mỗi thành phố, cho mọi mô hình và mốc. Khoảng loại trừ chung sáu giờ có tính bảo thủ đối với mốc ngắn hơn. Siêu tham số và họ mô hình được chọn bằng RMSE validation gộp, tính từ tổng bình phương phần dư gộp, không phải trung bình các RMSE fold. Mỗi họ được fit lại trên development rồi đánh giá trên cùng các dòng test [5].'),
('h','4 1 Metrics and interpretation','4 1 Chỉ số và cách diễn giải'),
('p','RMSE is the square root of the mean squared residual; MAE is the mean absolute residual. R² compares residual sum of squares with the target’s total centered sum of squares. Relative improvement is 100 × (RMSE_persistence − RMSE_model) / RMSE_persistence. A negative improvement means the model is worse. City-level scores complement pooled scores because pooled squared errors can be dominated by more variable cities. No confidence intervals or significance tests are reported; temporally adjacent errors are not assumed independent.',
'RMSE là căn bậc hai của trung bình bình phương phần dư; MAE là trung bình trị tuyệt đối phần dư. R² so sánh tổng bình phương phần dư với tổng bình phương độ lệch của nhãn quanh trung bình. Mức cải thiện tương đối là 100 × (RMSE_persistence − RMSE_model) / RMSE_persistence. Giá trị âm nghĩa là mô hình kém hơn. Chỉ số theo thành phố bổ sung cho chỉ số gộp vì thành phố biến động lớn có thể chi phối sai số bình phương. Báo cáo không đưa khoảng tin cậy hay kiểm định ý nghĩa thống kê; không giả định các sai số liền kề độc lập.'),
('p','The software also records classification counts above an experimental hourly cutoff of 35.4 µg/m³. This is an auxiliary diagnostic, not an AQI calculation or a validated health warning. Regression predictions are not clipped to zero; negative outputs are counted explicitly.',
'Phần mềm còn ghi số đếm phân loại trên ngưỡng thực nghiệm theo giờ 35,4 µg/m³. Đây là chỉ báo phụ, không phải phép tính AQI hay cảnh báo sức khỏe đã được xác nhận. Dự đoán hồi quy không bị cắt về không; đầu ra âm được đếm tường minh.'),
],
[
('h','5 Results','5 Kết quả'),
('p','Multiple linear regression is selected at all horizons. Its pooled validation RMSE is 3.349, 7.743 and 11.144 µg/m³ at one, three and six hours. Ridge selects alpha = 0 and ties with multiple linear regression; the deterministic family order resolves the tie. The experiment therefore shows no validation benefit from positive L2 regularization within the tested grid. Polynomial regression selects degree two at every horizon [5].',
'Hồi quy tuyến tính đa biến được chọn ở cả ba mốc. RMSE validation gộp lần lượt là 3,349; 7,743 và 11,144 µg/m³ tại một, ba và sáu giờ. Ridge chọn alpha = 0 và đồng hạng với hồi quy đa biến; thứ tự họ mô hình cố định giải quyết trường hợp bằng điểm. Do đó, thí nghiệm không cho thấy lợi ích validation của phạt L2 dương trong lưới đã thử. Hồi quy đa thức chọn bậc hai ở cả ba mốc [5].'),
('allmetrics',),
('p','On test, polynomial regression has a slightly lower six-hour RMSE (13.876) than the validation-selected model (13.974). This descriptive result is retained without changing the selected model: choosing again after viewing test would contaminate the model-selection protocol. The difference is not claimed to be statistically significant.',
'Trên test, hồi quy đa thức có RMSE sáu giờ thấp hơn nhẹ (13,876) so với mô hình đã chọn bằng validation (13,974). Kết quả mô tả này được giữ nguyên mà không đổi lựa chọn: chọn lại sau khi xem test sẽ làm sai quy trình lựa chọn mô hình. Chênh lệch không được diễn giải là có ý nghĩa thống kê.'),
('selected',),
],
[
('h','5 1 Differences across cities','5 1 Khác biệt giữa các thành phố'),
('citymetrics',),
('p','The selected model improves on persistence in Hanoi, Hai Phong and Ho Chi Minh City, but loses in Da Nang at every horizon. At six hours, Da Nang RMSE rises from 4.675 for persistence to 6.261 for multiple linear regression, an increase of 33.91%. For that city and horizon, none of the 34 target values above the experimental cutoff are detected. These results limit any claim of uniform usefulness across cities. A city-specific or baseline-fallback policy would require a new validation protocol rather than a decision made from this test set.',
'Mô hình được chọn tốt hơn persistence tại Hà Nội, Hải Phòng và TP. Hồ Chí Minh, nhưng kém ở Đà Nẵng tại mọi mốc. Tại sáu giờ, RMSE Đà Nẵng tăng từ 4,675 của persistence lên 6,261 của hồi quy đa biến, tăng 33,91%. Với thành phố và mốc này, cả 34 nhãn vượt ngưỡng thực nghiệm đều không được phát hiện. Kết quả giới hạn khẳng định về tính hữu ích đồng đều giữa các thành phố. Chính sách riêng theo thành phố hoặc tự chuyển về baseline cần một quy trình validation mới, không được quyết định từ tập test này.'),
('fig','danang_1h.png','Figure 1. Da Nang one-hour forecasts over the last seven test days. “Observed” denotes the API target, not an independent station measurement. Each point uses a new forecast origin.','Hình 1. Dự báo Đà Nẵng trước một giờ trong bảy ngày test cuối. “Observed” là nhãn API, không phải số đo trạm độc lập. Mỗi điểm dùng một thời điểm phát hành mới.'),
],
[
('h','6 Discussion and limitations','6 Thảo luận và giới hạn'),
('p','Close agreement in Figure 1 is not by itself evidence of leakage. Persistence also follows the target closely at one hour, indicating strong short-range continuity. The test simulates repeated issuance: observations up to each origin are available, including observations acquired earlier within the test period. Model parameters remain fixed after development fitting. It does not simulate a single forecast of the whole week made at its start. Using current-hour features is legitimate only under the stated availability assumption.',
'Các đường gần nhau trong Hình 1 tự nó không chứng minh rò rỉ dữ liệu. Persistence cũng bám sát nhãn ở mốc một giờ, thể hiện tính liên tục ngắn hạn mạnh. Test mô phỏng phát hành lặp lại: quan sát đến mỗi thời điểm phát hành được xem là đã có, bao gồm quan sát vừa nhận trong giai đoạn test. Tham số mô hình giữ nguyên sau khi fit development. Đây không phải dự báo cả tuần một lần từ đầu tuần. Dùng đặc trưng giờ hiện tại chỉ hợp lệ dưới giả định dữ liệu sẵn có đã nêu.'),
('p','The code constrains lags and rolling statistics to past and current observations, separates target columns from the feature allowlist, fits scalers on training data, and purges future labels at split boundaries. Automated tests cover these mechanisms. They reduce implementation leakage risk but cannot verify historical release times of upstream data. The provider documents a three-hour native temporal resolution for its CAMS global product [1]; an hourly API series should therefore not be interpreted as independent hourly station observations. Provider processing and historical weather revisions may make this retrospective task easier than live forecasting.',
'Mã nguồn giới hạn độ trễ và thống kê trượt ở quan sát quá khứ và hiện tại, tách nhãn khỏi danh sách đặc trưng, fit scaler trên train và loại nhãn vượt ranh giới. Kiểm thử tự động bao phủ các cơ chế này. Chúng giảm rủi ro rò rỉ do cài đặt nhưng không xác minh được thời điểm nhà cung cấp thực sự công bố dữ liệu trong quá khứ. Nhà cung cấp mô tả độ phân giải thời gian gốc ba giờ cho sản phẩm CAMS toàn cầu [1]; vì vậy chuỗi API theo giờ không nên được xem như các số đo trạm độc lập từng giờ. Xử lý phía nguồn và hiệu chỉnh thời tiết lịch sử có thể khiến bài toán hồi cứu dễ hơn dự báo trực tiếp.'),
('p','Other limitations include approximately six months of missing boundary-layer height, complete-case selection, one final seasonal test interval, no station-based external validation, no prediction intervals and no feature-ablation study. The selected model produces 0, 8 and 16 negative test predictions at one, three and six hours. Unconstrained linear regression can violate concentration non-negativity. Future work should first establish data availability and station-based evaluation, then test predeclared missing-data strategies, city-specific models and nonnegative or calibrated uncertainty outputs using new validation periods.',
'Các giới hạn khác gồm gần sáu tháng thiếu độ cao lớp biên, chọn mẫu đầy đủ, chỉ một khoảng test cuối mang tính mùa vụ, chưa kiểm chứng ngoài mẫu bằng trạm, chưa có khoảng dự báo và chưa có nghiên cứu loại từng nhóm đặc trưng. Mô hình được chọn tạo 0, 8 và 16 dự đoán test âm tại một, ba và sáu giờ. Hồi quy tuyến tính không ràng buộc có thể vi phạm tính không âm của nồng độ. Hướng tiếp theo nên xác lập thời điểm dữ liệu sẵn có và kiểm chứng bằng trạm, sau đó thử chiến lược dữ liệu thiếu được định trước, mô hình theo thành phố và đầu ra không âm hoặc bất định được hiệu chỉnh trên giai đoạn validation mới.'),
('h','7 Conclusion','7 Kết luận'),
('p','Direct regression yields a reproducible benchmark for three PM2.5 forecast horizons. Validation-selected multiple linear regression improves pooled test RMSE by approximately 18–20% over persistence, while the city-level analysis reveals consistent underperformance in Da Nang. Ridge at zero penalty matches OLS, and polynomial test performance does not justify revising the validation decision. The central contribution is an auditable comparison with explicit temporal assumptions and limitations, not proof of real-time readiness or general superiority across locations.',
'Hồi quy trực tiếp tạo một bộ đánh giá có thể tái lập cho ba mốc dự báo PM2.5. Hồi quy đa biến được chọn bằng validation giảm RMSE test gộp khoảng 18–20% so với persistence, nhưng phân tích theo thành phố cho thấy kém ổn định tại Đà Nẵng ở cả ba mốc. Ridge không phạt trùng kết quả OLS và điểm test đa thức không phải lý do đổi quyết định validation. Đóng góp chính là phép so sánh có thể kiểm tra với giả định thời gian và giới hạn rõ ràng, chưa phải bằng chứng sẵn sàng vận hành hay ưu thế chung cho mọi địa điểm.'),
],
[
('h','8 Reproducibility and references','8 Tái lập và tài liệu tham khảo'),
('p','All reported measurements come from outputs/runs/current [5]. The run stores its configuration, feature order, selected parameters, library versions, source hashes, raw-data hashes and per-row predictions. Its base Git commit is dc3ebc3fddf1e459741c378ed4e9aabae28a30d7, but the working tree is modified; reproducing the study requires the local source snapshot rather than that commit alone. The training environment used Python 3.10.11, NumPy 2.2.6, pandas 2.3.3 and scikit-learn 1.7.2. Twenty automated tests passed, including causal feature construction, boundary purging, train-only scaling and model save/reload inference.',
'Toàn bộ số liệu báo cáo lấy từ outputs/runs/current [5]. Run lưu cấu hình, thứ tự đặc trưng, tham số được chọn, phiên bản thư viện, mã băm mã nguồn và dữ liệu, cùng dự đoán từng dòng. Git commit nền là dc3ebc3fddf1e459741c378ed4e9aabae28a30d7 nhưng working tree đã thay đổi; tái lập cần snapshot mã nguồn local thay vì chỉ commit đó. Môi trường huấn luyện dùng Python 3.10.11, NumPy 2.2.6, pandas 2.3.3 và scikit-learn 1.7.2. Hai mươi kiểm thử tự động đạt, bao gồm đặc trưng nhân quả theo thời gian, loại trừ ranh giới, scaler chỉ fit train và lưu/nạp model để dự đoán.'),
('p','From the repository root, the following commands create a new run and issue forecasts. The destination must not already exist. The inference example uses the final historical origin, 31 December 2024 23:00, and produces twelve rows for 1 January 2025 at 00:00, 02:00 and 05:00. These are historical demonstration outputs, not forecasts for the report date.',
'Từ gốc repository, các lệnh dưới đây tạo run mới và phát hành dự đoán. Thư mục đích chưa được tồn tại. Ví dụ inference dùng thời điểm lịch sử cuối, 23:00 ngày 31/12/2024, tạo mười hai dòng cho 00:00, 02:00 và 05:00 ngày 01/01/2025. Đây là đầu ra minh họa lịch sử, không phải dự báo cho ngày viết báo cáo.'),
('code',),
('refs',),
]
]

def build(lang):
 d=Document(); sec=d.sections[0]; sec.page_height=Inches(11.69); sec.page_width=Inches(8.27)
 sec.top_margin=sec.bottom_margin=Inches(.72); sec.left_margin=sec.right_margin=Inches(.8)
 for name in ['Normal','Title','Subtitle','Heading 1','Caption']:
  st=d.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(11); d.styles['Normal'].paragraph_format.space_after=Pt(7); d.styles['Normal'].paragraph_format.line_spacing=1.08
 d.styles['Title'].font.size=Pt(21); d.styles['Subtitle'].font.size=Pt(12)
 d.styles['Heading 1'].font.size=Pt(13); d.styles['Heading 1'].paragraph_format.space_before=Pt(9)
 d.styles['Caption'].font.size=Pt(9)
 f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
 fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
 def p(text,style=None):
  x=d.add_paragraph(text,style); x.paragraph_format.widow_control=True; return x
 def table(headers,rows):
  t=d.add_table(rows=1, cols=len(headers)); t.autofit=False
  widths=[2.3]+[(6.67-2.3)/(len(headers)-1)]*(len(headers)-1)
  for c,w in zip(t.columns,widths):c.width=Inches(w)
  for c,s in zip(t.rows[0].cells,headers):c.text=s
  for row in rows:
   for c,s in zip(t.add_row().cells,row):c.text=str(s)
  for i,row in enumerate(t.rows):
   pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
   if i==0:pr.append(OxmlElement('w:tblHeader'))
   for j,c in enumerate(row.cells):
    for pp in c.paragraphs:
     pp.paragraph_format.space_after=Pt(4); pp.paragraph_format.space_before=Pt(4)
     for r in pp.runs:r.font.size=Pt(9); r.bold=i==0
    tcp=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    tcp.append(borders)
    if i==0:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF3');tcp.append(sh)
  p('')
 en=lang=='EN'; ix=1 if en else 2
 for n,page in enumerate(pages):
  if n:d.add_page_break()
  for item in page:
   typ=item[0]
   if typ in ['title','sub','h','p']:
    p(item[ix],{'title':'Title','sub':'Subtitle','h':'Heading 1','p':None}[typ])
   elif typ=='cities':
    table(['City' if en else 'Thành phố','Latitude','Longitude','Rows' if en else 'Số dòng'],[['Hanoi',21.0285,105.8542,17544],['Hai Phong',20.8449,106.6881,17544],['Da Nang',16.0544,108.2022,17544],['Ho Chi Minh City',10.8231,106.6297,17544]])
   elif typ=='folds':
    p('Table I. Expanding validation folds.' if en else 'Bảng I. Các fold validation mở rộng.','Caption')
    table(['Fold','Train rows' if en else 'Dòng train','Validation rows' if en else 'Dòng validation'],[[1,36380,2576],[2,38980,2576],[3,41580,2576]])
   elif typ=='allmetrics':
    p('Table II. Held-out test RMSE in µg/m³, 7,804 rows per horizon.' if en else 'Bảng II. RMSE test (µg/m³), 7.804 dòng mỗi mốc.','Caption')
    rows=[]
    for family in ['persistence','simple_linear','multiple_linear','polynomial','ridge']:
     rows.append([family]+[f"{float(next(r for r in metrics if r['model']==family and r['horizon_hours']==str(h))['rmse']):.3f}" for h in [1,3,6]])
    table(['Model' if en else 'Mô hình','+1 h','+3 h','+6 h'],rows)
   elif typ=='selected':
    p('Table III. Selected multiple linear model on test.' if en else 'Bảng III. Mô hình hồi quy đa biến được chọn trên test.','Caption')
    table(['Horizon' if en else 'Mốc','MAE','R²','Gain %' if en else 'Cải thiện %'],[[r['horizon_hours']+' h',f"{float(r['mae']):.3f}",f"{float(r['r2']):.3f}",f"{float(r['improvement_pct']):.2f}"] for r in metrics if r['model']=='multiple_linear'])
   elif typ=='citymetrics':
    p('Table IV. Selected model RMSE and improvement over persistence by city.' if en else 'Bảng IV. RMSE mô hình được chọn và cải thiện so với persistence theo thành phố.','Caption')
    table(['City / h' if en else 'Thành phố / giờ','RMSE','Baseline','Gain %' if en else 'Cải thiện %'],[[r['city']+' / '+r['horizon_hours'],f"{float(r['rmse']):.3f}",f"{float(r['baseline_rmse']):.3f}",f"{float(r['improvement_pct']):.2f}"] for r in sorted([r for r in city if r['model']=='multiple_linear'],key=lambda r:(r['city'],int(r['horizon_hours'])))])
   elif typ=='fig':
    d.add_picture(str(RUN/'figures'/item[1]),width=Inches(6.5));p(item[2 if en else 3],'Caption')
   elif typ=='code':
    for line in ['python3 air_quality_forecasting/main.py --skip-fetch \\', '  --run-dir air_quality_forecasting/outputs/runs/reproduce', 'python3 air_quality_forecasting/predict.py \\', '  --run-dir air_quality_forecasting/outputs/runs/reproduce \\', '  --history-dir air_quality_forecasting/data/raw_api']:
     pp=p(line);pp.paragraph_format.space_after=Pt(0)
     for r in pp.runs:r.font.name='Courier New';r.font.size=Pt(8)
   elif typ=='refs':
    p('References' if en else 'Tài liệu tham khảo','Heading 1')
    for i,ref in enumerate(refs,1):
     pp=p(f'[{i}] {ref}')
     for r in pp.runs:r.font.size=Pt(9)
 d.core_properties.title=pages[0][0][ix]; d.core_properties.author='';d.core_properties.subject='PM2.5 regression forecasting research report'
 dest=OUT/f'PM25_Research_Report_{lang}.docx';d.save(dest);print(dest)
for lang in ['EN','VI']:build(lang)
